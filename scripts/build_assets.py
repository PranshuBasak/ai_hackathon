"""Generate fictional, deterministic local fixtures. Does not access Creatio."""
from pathlib import Path
import csv, json, uuid
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'seed-data'
OUT.mkdir(exist_ok=True)
NS = uuid.UUID('ec9d08cf-6716-42e6-b048-e1858ba06b56')
def uid(key): return str(uuid.uuid5(NS, key))
def write_csv(name, rows):
    with (OUT/name).open('w', encoding='utf-8-sig', newline='') as f:
        w=csv.DictWriter(f, fieldnames=list(rows[0])); w.writeheader(); w.writerows(rows)

groups = {
'Developer':['Harborline Development Group','Summit Ridge Properties','Cedar & Stone Developments','Northgate Urban Partners','Bluewater Living Co.'],
'Architect':['Arcline Architects','Meridian Design Studio','Halvorsen + Reyes Architecture','Studio Lumen'],
'Contractor':['Ironclad Builders','Keystone Construction Co.','Pinnacle General Contracting'],
'Dealer':['Metro Appliance Supply','Coastal Kitchen Distributors','Prairie Home Appliance','Lakeshore Builder Supply'],
'Our company':['Everline Appliances'],
'Partner':['Crest Interior Studio','Copperleaf Engineering','Vista Specification Partners']}
accounts=[]
for typ,names in groups.items():
    for i,name in enumerate(names):
        accounts.append(dict(Id=uid('account:'+name),Name=name,Type=typ,AccountCategory=(['A','B','C','B'][i] if typ=='Dealer' else ''),AlternativeName=('Harborline Dev Group; HDG' if name.startswith('Harborline') else ''),SeedTag='HackathonSeed'))
contacts=[]
for i,name in enumerate(['Avery Lane','Jordan Reed','Morgan Brooks','Taylor Quinn','Casey Blake','Riley Hart','Jamie Ellis','Cameron Vale','Alex Rowan','Drew Parker','Robin Hayes','Sam Emery','Quinn Avery','Reese Harper','Skyler Gray']):
    account=(groups['Developer']+groups['Architect']+groups['Dealer'])[i%13]
    contacts.append(dict(Id=uid('contact:'+name),Name=name,Account=account,Email=f'fixture{i+1:02d}@example.com',JobTitle='Project liaison',SeedTag='HackathonSeed'))
names=['The Meridian Tower – Phase 1','Riverside Commons','Pier 9 Residences','Cedar Grove Heights','Northgate Hotel','Bluewater Gardens','Summit Student Village','Lumen Senior Living','Copperleaf Residences','Stonebridge Commons','Willow Creek Apartments','Lakeview Terrace','Vista Mixed Use','Orchard Square','Maple Court','Juniper Heights','Prairie Village','Harbor Point Hotel','Canyon Apartments','Elm Street Residences','Pinecrest Living','Westgate Place','Clearwater Court','Oakline Campus','Meadowbrook Community']
cities=['Austin','Denver','Miami','Phoenix','Seattle','Chicago','Charlotte','Nashville']
states=['TX','CO','FL','AZ','WA','IL','NC','TN']
types=['Multifamily High-Rise','Multifamily Mid-Rise','Mixed-Use','Hospitality','Senior Living','Student Housing','Single-Family Community']
stages=['Conceptual','Design development','Construction documents','Bidding','Under construction','Completed']
projects=[]; participants=[]; opportunities=[]
for i,name in enumerate(names):
    pid=uid('project:'+name); dev=groups['Developer'][i%5]
    projects.append(dict(Id=pid,UsrProjectName=name,Type=types[i%7],Status='Planned',ProjectEntryType='RESOLVE_DEFAULT_AT_LOAD',Owner='RESOLVE_CURRENT_CONTACT_AT_LOAD',Account=dev,UsrExternalId=f'BASE-{i+1:03d}',UsrExternalSource='Dodge',UsrCity=cities[i%8],UsrProjectAddress=f'{100+i*10} Example Avenue',UsrProjectCountry='United States',UsrConstructionStage=stages[i%6],UsrEstimatedProjectValue=40000000+i*5000000,UsrNumberOfUnits=100+i*10,UsrPriorityClassification=('Strategic Pursuit' if i%5==0 else 'Active pursuit'),ParentProject='',Opportunity=(uid('opportunity:'+name) if i<8 else ''),SeedTag='HackathonSeed'))
    roles=[('Developer/Owner',dev),('General contractor/Builder',groups['Contractor'][i%3]),('Dealer',groups['Dealer'][i%4])]
    if i<20: roles.append(('Architect/Specifier',groups['Architect'][i%4]))
    for role,account in roles:
        participants.append(dict(Id=uid('party:'+name+role),UsrProject=pid,UsrAccount=account,UsrContact='',UsrPartyRole=role,UsrIsPrimary=(role=='Developer/Owner'),UsrSourceIntake='',SeedTag='HackathonSeed'))
    if i<8:
        opportunities.append(dict(Id=uid('opportunity:'+name),Title=name+' appliance pursuit',Account=dev,Partner=groups['Dealer'][i%4],Budget=round((40000000+i*5000000)*0.015),Stage='RESOLVE_EARLY_NON_PROPOSAL_STAGE',UsrADProject=pid,SeedTag='HackathonSeed'))
headers=['Source','Source Project ID','Project Name','Project Type','Stage','Address','City','State','Country','Est. Construction Value','Units','Bid Date','Start Date','Completion Date','Owner/Developer','Architect','GC/Builder','Dealer','Key Contact','Contact Email','Contact Role','Description']
batch=[]; oracle=[]
def row(n,name,city='Austin',value=65000000,units=160):
    return dict(zip(headers,['Dodge',f'PI-DEMO-{n:02d}',name,'Multifamily High-Rise','Design development',f'{700+n*10} Example Avenue',city,'TX','United States',value,units,'2026-11-15','2027-02-01','2029-06-30',groups['Developer'][n%5],groups['Architect'][n%4],groups['Contractor'][n%3],groups['Dealer'][n%4],contacts[n%15]['Name'],contacts[n%15]['Email'],'Project liaison',f'Fictional HackathonSeed intake scenario {n}.']))
for n in range(1,16):
    r=row(n,['Aurora Skyline','Elmwood Residences','Meadow Senior Living'][n-1] if n<=3 else f'Scenario {n}')
    action='Create new project'; priority='Active pursuit'; target=''; lifecycle='Ready to apply'; rule='High confidence and all mandatory information resolved'
    if n==1:r['Est. Construction Value']=180000000;r['Units']=320;priority='Strategic Pursuit'
    if n==3:r['Est. Construction Value']=22000000;r['Units']=80;priority='Monitor'
    if n in [4,5,6,7,8,13,14,15]:
        p=projects[{4:3,5:4,6:5,7:0,8:1,13:6,14:2,15:2}[n]];target=p['Id']
        r.update({'Project Name':p['UsrProjectName'],'City':p['UsrCity'],'State':states[cities.index(p['UsrCity'])],'Address':p['UsrProjectAddress'],'Owner/Developer':p['Account'],'Est. Construction Value':p['UsrEstimatedProjectValue'],'Units':p['UsrNumberOfUnits']})
        action='Duplicate – no action';priority=p['UsrPriorityClassification'];rule='Same project; never create another Project'
        if n==4:r['Source Project ID']=p['UsrExternalId']
    if n==7:r['Project Name']='The Meridian Tower – Phase 2';action='Link as new phase';priority='Active pursuit';lifecycle='Needs review';rule='Phase always needs review regardless of threshold'
    if n==8:r['Project Name']='The Commons at Riverside';action='Update existing project';priority='Active pursuit';lifecycle='Needs review';rule='Rename always needs review regardless of threshold'
    if n==9:r['Project Name']='Harborline West Residences';r['Owner/Developer']='Harbourline Developmnt Grp';lifecycle='Needs review';rule='Developer resolves to Harborline Development Group; reviewer confirms alias'
    if n==10:r['Architect']='';priority='Data Incomplete';lifecycle='Needs review';rule='Missing architect; future Apply creates follow-up task'
    if n==11:r['Dealer']='';priority='Monitor';lifecycle='Needs review';rule='Missing dealer; flag missing information'
    if n==12:
        r.update({'Project Name':'Small Home Remodel','Project Type':'Single-Family','Est. Construction Value':150000,'Units':1,'Owner/Developer':'','Architect':'','GC/Builder':'','Dealer':'','Key Contact':'','Contact Email':''});action='Discard (low value)';priority='Low priority';lifecycle='Needs review';rule='Low fit and missing parties; do not create pipeline'
    if n==13:r['Stage']='Bidding';r['Est. Construction Value']=175000000;action='Update existing project';priority='Strategic Pursuit'
    if n==14:priority='Active pursuit';action='Update existing project|Duplicate – no action'
    if n==15:r['Source']='Tender portal';r['Source Project ID']='TENDER-PIER9-2026';priority='Active pursuit'
    batch.append(r)
    oracle.append(dict(Row=str(n),Source=r['Source'],SourceProjectId=r['Source Project ID'],ProjectName=r['Project Name'],ExpectedMatchedProjectId=target,ExpectedAction=action,ExpectedPriority=priority,ExpectedLifecycle=lifecycle,Rule=rule,Result='NOT RUN'))
for name,rows in [('accounts.csv',accounts),('contacts.csv',contacts),('projects.csv',projects),('participants.csv',participants),('opportunities.csv',opportunities),('ProjectIntake_batch.csv',batch)]:write_csv(name,rows)
links=[dict(Id=uid('oppcontact:'+str(i)),Opportunity=opportunities[i]['Id'],Contact=contacts[i]['Id'],Role='Influencer',SeedTag='HackathonSeed') for i in range(8)]
write_csv('opportunity_contacts.csv',links)
with (ROOT/'tests/oracle.csv').open('w',encoding='utf-8-sig',newline='') as f:
    w=csv.DictWriter(f,fieldnames=list(oracle[0]));w.writeheader();w.writerows(oracle)
fields=['UsrExternalSource','UsrExternalProjectId','UsrProjectName','UsrProjectTypeText','UsrStageText','UsrProjectAddress','UsrCityText','UsrStateText','UsrCountryText','UsrEstimatedProjectValue','UsrNumberOfUnits','UsrBidDate','UsrStartDate','UsrCompletionDate','UsrDeveloperName','UsrArchitectName','UsrBuilderName','UsrDealerName','UsrKeyContactName','UsrKeyContactEmail','UsrKeyContactRoleText','UsrProjectDescription']
mapping=[dict(header=h,field=f,match_key=f in ['UsrExternalSource','UsrExternalProjectId']) for h,f in zip(headers,fields)]
(ROOT/'docs/import-mapping.json').write_text(json.dumps(dict(entity='UsrADProjectIntelligence',mapping=mapping,excluded=['UsrStatus','UsrReviewedBy','UsrMatchedProject','UsrQualificationScore','UsrPriorityClassification','UsrRecommendedAction','UsrAnalysisDate','UsrAgentError'],platform_template_saved=False),indent=2),encoding='utf-8')
wb=Workbook();ws=wb.active;ws.title='Intake template';ws.append(headers)
batchws=wb.create_sheet('15 row batch');batchws.append(headers)
for r in batch:batchws.append(list(r.values()))
mapws=wb.create_sheet('Mapping');mapws.append(['Header','Entity field','Match key'])
for r in mapping:mapws.append([r['header'],r['field'],r['match_key']])
for sheet in wb:
    sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
    for cell in sheet[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='1F4E78')
    for col in sheet.columns:sheet.column_dimensions[col[0].column_letter].width=min(42,max(18,max(len(str(c.value or '')) for c in col)+2))
wb.save(OUT/'ProjectIntake_Import_Template.xlsx')
samples={
'email_E1.txt':'From: Avery Lane <architect@example.com>\nSubject: [INTAKE] Bayview Senior Living\nPlease log Bayview Senior Living in Tampa, Florida, US: 210 units, estimated construction value $85 million. Construction documents are due November 2026. Developer: Bluewater Living Co. Architect: Arcline Architects. Builder: Ironclad Builders. Dealer not yet selected. Contact Avery Lane, architect@example.com, Architect.\n',
'sample_texts/E2.txt':'Oak Terrace Hotel, Nashville TN, United States. Hospitality, 140 units, $60 million, Design development. Developer Summit Ridge Properties; architect Studio Lumen; builder Keystone Construction Co.; dealer Metro Appliance Supply. Contact Jordan Reed at jordan@example.com.\n',
'sample_texts/E3.txt':'The Meridian Tower – Phase 2 at 100 Example Avenue, Austin TX, United States. 240 apartments, $120 million, Construction documents. Harborline Development Group owns it, Arcline Architects specifies it, Ironclad Builders is GC, Metro Appliance Supply is dealer. This is a new phase, not a duplicate.\n',
'sample_texts/E4.txt':'Forwarded project note: The Commons at Riverside is the new name for Riverside Commons, 110 Example Avenue, Denver CO, US. Developer Summit Ridge Properties. $45 million, 110 units, Bidding. Architect and dealer not supplied.\n'}
for name,content in samples.items():(OUT/name).parent.mkdir(exist_ok=True);(OUT/name).write_text(content,encoding='utf-8')
extractions=[dict(sample='E1',projectName='Bayview Senior Living',city='Tampa',state='Florida',country='United States',units=210,constructionValue=85000000,stage='Construction documents',developer='Bluewater Living Co.',architect='Arcline Architects',builder='Ironclad Builders',dealer=None,keyContact={'name':'Avery Lane','email':'architect@example.com','role':'Architect'},bidDate=None,dateNote='November 2026 only; do not invent a day'),dict(sample='E2',projectName='Oak Terrace Hotel',city='Nashville',state='TN',country='United States',units=140,constructionValue=60000000,stage='Design development',developer='Summit Ridge Properties',architect='Studio Lumen',builder='Keystone Construction Co.',dealer='Metro Appliance Supply'),dict(sample='E3',projectName='The Meridian Tower – Phase 2',city='Austin',units=240,constructionValue=120000000,stage='Construction documents',matchType='new-phase',status='Needs review'),dict(sample='E4',projectName='The Commons at Riverside',city='Denver',units=110,constructionValue=45000000,stage='Bidding',matchType='renamed',status='Needs review',architect=None,dealer=None)]
(ROOT/'tests/extraction-oracle.json').write_text(json.dumps(extractions,indent=2),encoding='utf-8')
fixtures={entity:[dict(id=r['Id'],key=r.get('Name',r.get('UsrProjectName',r.get('Title',r['Id']))),loaded=False) for r in rows] for entity,rows in [('Account',accounts),('Contact',contacts),('Project',projects),('UsrADProjectParty',participants),('Opportunity',opportunities),('OpportunityContact',links)]}
(OUT/'fixture-manifest.json').write_text(json.dumps(dict(version=1,environment='ai_hackathon',tag='HackathonSeed',state='PLANNED_NOT_LOADED',baseline=fixtures,intakeKeys=[{'source':r['Source'],'sourceProjectId':r['Source Project ID'],'id':None,'loaded':False} for r in batch],runRecords=[],resetRule='Delete only verified loaded GUIDs in this manifest; null IDs are never eligible'),indent=2),encoding='utf-8')
assert len(accounts)==20 and len(contacts)==15 and len(projects)==25 and len(opportunities)==8 and len(batch)==15
assert len({(r['Source'],r['Source Project ID']) for r in batch})==15
assert sum(1 for p in projects if not any(x['UsrProject']==p['Id'] and x['UsrPartyRole']=='Architect/Specifier' for x in participants))==5
assert all(not r['ExpectedMatchedProjectId'] or any(p['Id']==r['ExpectedMatchedProjectId'] for p in projects) for r in oracle)
print('Assets validated: 20 accounts, 15 contacts, 25 projects, 95 participants, 8 opportunities, 8 opportunity contacts, 15 intakes, 4 text samples.')
