"""Create local-only fictional input for the Project Intake Capture Assistant."""
import csv
import json
from datetime import date
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

root = Path(__file__).resolve().parents[1]
headers = [m['header'] for m in json.loads((root/'docs/import-mapping.json').read_text())['mapping']]
rows = [
 ['Dodge','CAPTEST-20260927-001','Oakline Court Residences','Multifamily High-Rise','Design development','101 Fictional Oakline Avenue','Austin','TX','United States',48000000,120,date(2026,12,10),date(2027,3,1),date(2028,9,30),'Bluewater Living Co.','Arcline Architects','Ironclad Builders','Coastal Kitchen Distributors','Avery Lane','oakline.test@example.com','Project liaison','Fictional capture test 1: complete source information. Construction value is explicitly USD 48,000,000; not expected product order value.'],
 ['Dodge','CAPTEST-20260927-002','Juniper Harbor Apartments','Apartments','Construction documents','202 Fictional Juniper Street','Seattle','WA','United States',72000000,180,date(2027,1,15),date(2027,5,1),date(2029,2,28),'Summit Ridge Properties','Meridian Design Studio','Keystone Construction Co.','Coastal Kitchen Distributors','Jordan Reed','juniper.test@example.com','Development manager','Fictional capture test 2: complete source information. Construction value is explicitly USD 72,000,000.'],
 ['Tender portal','CAPTEST-20260927-003','Silverleaf Senior Living','Senior Living','Bidding','303 Fictional Silverleaf Road','Denver','CO','United States',36000000,90,date(2026,11,20),date(2027,2,1),date(2028,8,31),'Cedar & Stone Developments','Studio Lumen','Pinnacle General Contracting','Prairie Home Appliance','Morgan Brooks','silverleaf.test@example.com','Procurement contact','Fictional capture test 3: complete source information. Construction value is explicitly USD 36,000,000.'],
 ['Dodge','CAPTEST-20260927-004','Cedarbrook Student Village','Student Housing','Design development','404 Fictional Cedarbrook Lane','Nashville','TN','United States',55000000,240,date(2027,2,10),date(2027,6,1),date(2029,5,31),'Northgate Urban Partners','','Ironclad Builders','','Taylor Quinn','cedarbrook.test@example.com','Project coordinator','Fictional capture test 4. Construction value is explicitly USD 55,000,000. Architect and dealer have not been selected; keep those fields empty.'],
 ['Tender portal','CAPTEST-20260927-005','Willowgate Mixed-Use Development','Mixed-Use','Conceptual','505 Fictional Willowgate Boulevard','Charlotte','NC','United States',None,None,None,None,None,'Harborline Development Group','Arcline Architects','','','Casey Morgan','willowgate.test@example.com','Development coordinator','Fictional capture test 5: early concept. Value, units, bid/start/completion dates, builder and dealer are unknown. Design review is expected in November 2026; this is not a project completion date and no exact day has been given.'],
]
assert len(headers)==22 and all(len(r)==22 for r in rows)
target=root/'seed-data/capture-assistant-test'
target.mkdir(exist_ok=True)
wb=Workbook(); ws=wb.active; ws.title='Project Intakes'
ws.append(headers)
for row in rows: ws.append(row)
ws.freeze_panes='D2'; ws.sheet_view.zoomScale=80
for c in ws[1]:
    c.fill=PatternFill('solid',fgColor='16324F'); c.font=Font(color='FFFFFF',bold=True); c.alignment=Alignment(wrap_text=True,vertical='center')
ws.row_dimensions[1].height=34
for row in ws.iter_rows(min_row=2):
    ws.row_dimensions[row[0].row].height=78
    for c in row: c.alignment=Alignment(wrap_text=True,vertical='top')
    row[9].number_format='#,##0'; row[10].number_format='0'
    for c in row[11:14]: c.number_format='yyyy-mm-dd'
for i in range(1,23): ws.column_dimensions[get_column_letter(i)].width=24
for letter,width in {'A':17,'B':27,'C':36,'F':34,'J':23,'K':12,'L':16,'M':16,'N':18,'T':32,'V':65}.items(): ws.column_dimensions[letter].width=width
table=Table(displayName='CaptureTestInputs',ref='A1:V6'); table.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True); ws.add_table(table)
ws.sheet_properties.pageSetUpPr.fitToPage=True
ws.page_setup.orientation='landscape'; ws.page_setup.paperSize=ws.PAPERSIZE_A3; ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=1
ws.print_title_rows='1:1'; ws.print_options.horizontalCentered=True
out=target/'Project_Intake_Capture_Test_5.xlsx'; wb.save(out)
with (target/'Project_Intake_Capture_Test_5.csv').open('w',encoding='utf-8-sig',newline='') as f:
    w=csv.writer(f); w.writerow(headers); w.writerows(rows)
check=load_workbook(out,data_only=True); sh=check['Project Intakes']
assert sh.max_row==6 and sh.max_column==22
assert [c.value for c in sh[1]]==headers
assert len({sh.cell(i,2).value for i in range(2,7)})==5
assert sh['P5'].value is None and sh['R5'].value is None
assert all(sh.cell(6,c).value is None for c in range(10,15))
assert not any(c.data_type=='f' for row in sh for c in row)
existing=list(csv.DictReader((root/'seed-data/ProjectIntake_batch.csv').open(encoding='utf-8-sig')))
assert not {r[1] for r in rows}.intersection(r['Source Project ID'] for r in existing)
print(f'Created and verified {out}: 5 records, 22 exact headers, real Excel dates/numbers, unique keys, intentional blanks, no formulas. No platform writes.')
