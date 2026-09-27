from pathlib import Path
import json,re,csv,hashlib
from openpyxl import load_workbook
R=Path(__file__).resolve().parents[1]
d=json.loads((R/'evidence/2026-09-26-data-model.json').read_text(encoding='utf-8'))
p=R/'docs/02-data-model.md'; old=p.read_text(encoding='utf-8')
head=old.split('## UsrADProjectIntelligence')[0]
tail='## Default and preservation rules'+old.split('## Default and preservation rules')[1]
blocks=[]
for entity in d['entities']:
    if not entity['columns']:continue
    lines=['## '+entity['name'],'','Live package-owned column read-back, 2026-09-26. Inherited columns are available through effective metadata.','','| Column | Caption | Type | Reference | Required | Default |','|---|---|---|---|---|---|']
    for c in sorted(entity['columns'],key=lambda c:c['name']):
        default=json.dumps(c.get('default-value-config',{}),ensure_ascii=False)
        lines.append('| '+' | '.join([c['name'],c.get('caption',''),c.get('type',''),c.get('reference-schema-name',''),str(c.get('required',False)).lower(),default])+' |')
    blocks.append('\n'.join(lines))
p.write_text(head+'\n\n'.join(blocks)+'\n\n'+tail,encoding='utf-8')
report={'verifiedAt':'2026-09-26','counts':{},'checks':{}}
for n,count in [('accounts',20),('contacts',15),('projects',25),('participants',95),('opportunities',8),('ProjectIntake_batch',15)]:
    rows=list(csv.DictReader((R/'seed-data'/f'{n}.csv').open(encoding='utf-8-sig')));assert len(rows)==count
    report['counts'][n]=len(rows)
batch=list(csv.DictReader((R/'seed-data/ProjectIntake_batch.csv').open(encoding='utf-8-sig')))
report['batchHeaders']=list(batch[0])
wb=load_workbook(R/'seed-data/ProjectIntake_Import_Template.xlsx',read_only=True,data_only=True)
report['workbookSheets']={s.title:{'rows':s.max_row,'columns':s.max_column} for s in wb}
source=Path(r'C:\Users\prans\.codex\attachments\04885667-333d-4e06-94cc-0b437618fdb0\Pasted text.txt')
assert source.read_bytes()==(R/'reference/original-plan.txt').read_bytes()
report['checks']['originalPlanByteIdentical']=True
report['originalPlanSha256']=hashlib.sha256(source.read_bytes()).hexdigest()
manifest=json.loads((R/'seed-data/fixture-manifest.json').read_text())
assert all(not x['loaded'] for arr in manifest['baseline'].values() for x in arr)
report['checks']['fixtureManifestNotLoaded']=True
(R/'evidence/static-assets-validation.json').write_text(json.dumps(report,indent=2),encoding='utf-8')
print(json.dumps(report))
